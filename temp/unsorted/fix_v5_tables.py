"""
Fix ICC_Companion_v5.docx - Add 6 missing database tables.
"""
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

V5_PATH = r'C:\xampp\htdocs\ICC_Companion\temp\ppt_and_report\my_report\ICC_Companion_v5.docx'
doc = Document(V5_PATH)

def make_table_element(headers, data):
    cols = len(headers)
    tbl = OxmlElement('w:tbl')
    tblPr = OxmlElement('w:tblPr')
    tblStyle = OxmlElement('w:tblStyle')
    tblStyle.set(qn('w:val'), 'Table Grid')
    tblPr.append(tblStyle)
    tblWidth = OxmlElement('w:tblW')
    tblWidth.set(qn('w:w'), '5000')
    tblWidth.set(qn('w:type'), 'pct')
    tblPr.append(tblWidth)
    tbl.append(tblPr)
    tblGrid = OxmlElement('w:tblGrid')
    for _ in range(cols):
        gridCol = OxmlElement('w:gridCol')
        gridCol.set(qn('w:w'), '2500')
        tblGrid.append(gridCol)
    tbl.append(tblGrid)

    def make_cell(text, is_header=False):
        tc = OxmlElement('w:tc')
        tcPr = OxmlElement('w:tcPr')
        tcW = OxmlElement('w:tcW')
        tcW.set(qn('w:w'), '2500')
        tcW.set(qn('w:type'), 'dxa')
        tcPr.append(tcW)
        if is_header:
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), 'D9E2F3')
            shading.set(qn('w:val'), 'clear')
            tcPr.append(shading)
        tc.append(tcPr)
        p = OxmlElement('w:p')
        pPr = OxmlElement('w:pPr')
        pJc = OxmlElement('w:jc')
        pJc.set(qn('w:val'), 'center')
        pPr.append(pJc)
        p.append(pPr)
        r = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), 'Times New Roman')
        rFonts.set(qn('w:hAnsi'), 'Times New Roman')
        rPr.append(rFonts)
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '22')
        rPr.append(sz)
        if is_header:
            b = OxmlElement('w:b')
            rPr.append(b)
        r.append(rPr)
        t = OxmlElement('w:t')
        t.text = str(text)
        t.set(qn('xml:space'), 'preserve')
        r.append(t)
        p.append(r)
        tc.append(p)
        return tc

    tr = OxmlElement('w:tr')
    for h in headers:
        tr.append(make_cell(h, is_header=True))
    tbl.append(tr)
    for row_data in data:
        tr = OxmlElement('w:tr')
        for val in row_data:
            tr.append(make_cell(val))
        tbl.append(tr)
    return tbl

def make_para_element(text, bold=True, size=12):
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    pJc = OxmlElement('w:jc')
    pJc.set(qn('w:val'), 'left')
    pPr.append(pJc)
    p.append(pPr)
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rPr.append(rFonts)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(size * 2))
    rPr.append(sz)
    if bold:
        b = OxmlElement('w:b')
        rPr.append(b)
    r.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    r.append(t)
    p.append(r)
    return p

def make_spacer():
    sp = OxmlElement('w:p')
    spPr = OxmlElement('w:pPr')
    sp.append(spPr)
    return sp

def insert_block_after(para_elem, heading_text, headers, data):
    """Insert heading + spacer + table + spacer after a paragraph element."""
    heading = make_para_element(heading_text, bold=True, size=12)
    table = make_table_element(headers, data)
    sp1 = make_spacer()
    sp2 = make_spacer()
    # Use addnext to append in order: heading, sp1, table, sp2
    # addnext inserts right after the element.
    # To get order: para -> heading -> sp1 -> table -> sp2:
    para_elem.addnext(sp2)
    para_elem.addnext(table)
    para_elem.addnext(sp1)
    para_elem.addnext(heading)

def insert_block_before(ref_elem, heading_text, headers, data):
    """Insert heading + spacer + table + spacer before a reference element."""
    heading = make_para_element(heading_text, bold=True, size=12)
    table = make_table_element(headers, data)
    sp1 = make_spacer()
    sp2 = make_spacer()
    # addprevious inserts elements in the order given.
    # Insert in desired order: heading -> sp1 -> table -> sp2 -> ref
    ref_elem.addprevious(heading)
    ref_elem.addprevious(sp1)
    ref_elem.addprevious(table)
    ref_elem.addprevious(sp2)

# ===== INSERT 4 TABLES AFTER THEIR PARAGRAPH HEADERS =====

# Find paragraphs by text content
def find_para(text_contains):
    for i, p in enumerate(doc.paragraphs):
        if text_contains in p.text:
            return i, p
    return None, None

# 1. announcements
_, para = find_para('Table 3: announcements')
if para:
    insert_block_after(para._element, '', 
        ['Field Name', 'Data Type', 'Constraint', 'Description'],
        [
            ('announcement_id', 'INT', 'PK, AUTO_INCREMENT', 'Unique announcement ID'),
            ('title', 'VARCHAR(255)', 'NOT NULL', 'Announcement title'),
            ('content', 'TEXT', 'NOT NULL', 'Announcement content'),
            ('priority', "ENUM('High','Medium','Low')", "DEFAULT 'Medium'", 'Priority level'),
            ('target_dept', 'VARCHAR(10)', 'NULL', 'Target department (null = all)'),
            ('target_semester', 'INT', 'NULL', 'Target semester (null = all)'),
            ('image_url', 'VARCHAR(255)', 'NULL', 'Announcement image'),
            ('created_at', 'TIMESTAMP', 'DEFAULT CURRENT_TIMESTAMP', 'Posting timestamp'),
        ]
    )
    print('Added announcements table')

# 2. class_routines
_, para = find_para('Table 4: class_routines')
if para:
    insert_block_after(para._element, '',
        ['Field Name', 'Data Type', 'Constraint', 'Description'],
        [
            ('routine_id', 'INT', 'PK, AUTO_INCREMENT', 'Unique routine ID'),
            ('dept_code', 'VARCHAR(10)', 'FK -> departments(dept_code)', 'Department code'),
            ('semester', 'INT', 'NOT NULL', 'Semester number'),
            ('file_url', 'VARCHAR(255)', 'NOT NULL', 'Routine file path (PDF/image)'),
            ('created_at', 'TIMESTAMP', 'DEFAULT CURRENT_TIMESTAMP', 'Upload timestamp'),
        ]
    )
    print('Added class_routines table')

# 3. resources
_, para = find_para('Table 5: resources')
if para:
    insert_block_after(para._element, '',
        ['Field Name', 'Data Type', 'Constraint', 'Description'],
        [
            ('resource_id', 'INT', 'PK, AUTO_INCREMENT', 'Unique resource ID'),
            ('title', 'VARCHAR(255)', 'NOT NULL', 'Resource title'),
            ('resource_type', "ENUM('syllabus','book','material','link','text','others')", 'NOT NULL', 'Type of resource'),
            ('dept_code', 'VARCHAR(10)', 'FK -> departments(dept_code)', 'Department code'),
            ('subject_id', 'INT', 'FK -> subjects(subject_id)', 'Linked subject (optional)'),
            ('file_url', 'VARCHAR(255)', 'NULL', 'File path or external URL'),
            ('description', 'TEXT', 'NULL', 'Resource description'),
            ('created_at', 'TIMESTAMP', 'DEFAULT CURRENT_TIMESTAMP', 'Upload timestamp'),
        ]
    )
    print('Added resources table')

# 4. lost_found
_, para = find_para('Table 6: lost_found')
if para:
    insert_block_after(para._element, '',
        ['Field Name', 'Data Type', 'Constraint', 'Description'],
        [
            ('item_id', 'INT', 'PK, AUTO_INCREMENT', 'Unique item ID'),
            ('title', 'VARCHAR(255)', 'NOT NULL', 'Item title'),
            ('description', 'TEXT', 'NULL', 'Item description'),
            ('item_type', "ENUM('Lost','Found')", 'NOT NULL', 'Lost or Found'),
            ('status', "VARCHAR(20)", "DEFAULT 'Pending'", 'Pending / Claimed / Returned'),
            ('category', 'VARCHAR(50)', 'NULL', 'Item category'),
            ('location', 'VARCHAR(100)', 'NULL', 'Where item was lost/found'),
            ('reported_by', 'VARCHAR(100)', 'NULL', 'Name of reporter'),
            ('contact', 'VARCHAR(50)', 'NULL', 'Contact info of reporter'),
            ('image_url', 'VARCHAR(255)', 'NULL', 'Item image path'),
            ('created_at', 'TIMESTAMP', 'DEFAULT CURRENT_TIMESTAMP', 'Submission timestamp'),
        ]
    )
    print('Added lost_found table')

# ===== INSERT exam_schedules (Table 7) AND faculty_subjects (Table 8) BEFORE CHAPTER 8 =====
_, ch8_para = find_para('CHAPTER 8')
if ch8_para:
    ch8_elem = ch8_para._element

    # Insert exam_schedules first (will appear first)
    insert_block_before(ch8_elem, 'Table 7: exam_schedules (Academic Calendar Scheduling)',
        ['Field Name', 'Data Type', 'Constraint', 'Description'],
        [
            ('schedule_id', 'INT', 'PK, AUTO_INCREMENT', 'Unique schedule ID'),
            ('dept_code', 'VARCHAR(10)', 'FK -> departments(dept_code)', 'Department code'),
            ('semester', 'INT', 'NOT NULL', 'Semester number'),
            ('schedule_type', "VARCHAR(20)", "NOT NULL", 'Sessional / Final / Comprehensive'),
            ('file_url', 'VARCHAR(255)', 'NOT NULL', 'Schedule file path (PDF/image)'),
            ('created_at', 'TIMESTAMP', 'DEFAULT CURRENT_TIMESTAMP', 'Upload timestamp'),
        ]
    )
    print('Added exam_schedules table')

    # Insert faculty_subjects next (will appear after exam_schedules)
    insert_block_before(ch8_elem, 'Table 8: faculty_subjects (Faculty Subject Assignments)',
        ['Field Name', 'Data Type', 'Constraint', 'Description'],
        [
            ('assignment_id', 'INT', 'PK, AUTO_INCREMENT', 'Unique assignment ID'),
            ('admin_id', 'INT', 'FK -> admins(admin_id)', 'Faculty/admin assigned'),
            ('subject_id', 'INT', 'FK -> subjects(subject_id)', 'Subject assigned'),
            ('created_at', 'TIMESTAMP', 'DEFAULT CURRENT_TIMESTAMP', 'Assignment timestamp'),
        ]
    )
    print('Added faculty_subjects table')

# ===== UPDATE DATABASE TABLE COUNT =====
for p in doc.paragraphs:
    if 'eleven' in p.text:
        for run in p.runs:
            if 'eleven' in run.text:
                run.text = run.text.replace('eleven', 'twelve')
                print('Updated table count to 12')
                break
        break

# ===== SAVE =====
doc.save(V5_PATH)
print(f'\nSaved: {V5_PATH}')

# ===== VERIFY =====
doc2 = Document(V5_PATH)
print(f'Total tables: {len(doc2.tables)}')
ft_count = 0
for t in doc2.tables:
    if t.rows and t.rows[0].cells[0].text.strip() == 'Field Name':
        ft_count += 1
print(f'Database schema tables: {ft_count}')
print(f'Other tables: {len(doc2.tables) - ft_count}')

# Show DB table listing
print('\nDatabase tables in order:')
for elem in doc2.element.body:
    tag = elem.tag.split('}')[-1]
    if tag == 'p':
        texts = [t.text for t in elem.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t') if t.text]
        txt = ''.join(texts).strip()
        if txt.startswith('Table ') and len(txt) < 100:
            print(f'  {txt}')
    elif tag == 'tbl':
        try:
            rows = elem.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr')
            if len(rows) > 1:
                tc = rows[1].find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')
                name = tc.text[:30] if tc is not None else '?'
                row_count = len(rows)
                print(f'  -> Table ({row_count}r): {name}')
            elif len(rows) == 1:
                print(f'  -> Table (header only)')
        except:
            print(f'  -> Table (error)')
