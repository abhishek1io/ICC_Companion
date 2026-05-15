"""Update student details in the existing final report."""
from docx import Document
import os, shutil

# Use the temp file (already has images)
src = os.path.join(os.environ.get('TEMP', 'C:/Temp'), 'ICC_FINAL_REPORT.docx')
if not os.path.exists(src):
    src = r'C:\xampp\htdocs\ICC_Companion\temp\ppt_and_report\my_report\PROJECT_REPORT.docx'

out = os.path.join(os.environ.get('TEMP', 'C:/Temp'), 'ICC_FINAL_REPORT.docx')

doc = Document(src)

def find_and_replace(search_text, replace_text):
    """Replace text in paragraphs only, preserving formatting."""
    count = 0
    for p in doc.paragraphs:
        if search_text in p.text:
            for r in p.runs:
                if search_text in r.text:
                    r.text = r.text.replace(search_text, replace_text)
                    count += 1
    return count

def replace_in_table(table, search, replace):
    count = 0
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if search in p.text:
                    for r in p.runs:
                        if search in r.text:
                            r.text = r.text.replace(search, replace)
                            count += 1
    return count

print("Replacing student details...")

# 1. Name changes
c = find_and_replace('Ankita Chetri', 'Abhishek Sharma')
print(f"  Name: {c} replacements")

c = find_and_replace('ANKITA CHETRI', 'ABHISHEK SHARMA')
print(f"  Name (upper): {c} replacements")

# 2. Roll number
c = find_and_replace('UT-231-049-0012', 'UT-231-049-0004')
print(f"  Roll: {c} replacements")

# 3. Registration number
c = find_and_replace('23084898', '23084891')
print(f"  Reg: {c} replacements")

# 4. Guide name
c = find_and_replace('Urbimala Hazarika', 'Manas Chakraborty')
print(f"  Guide: {c} replacements")

# 5. Pronouns (certificate + acknowledgement)
# "carried out by her" -> "carried out by him"
c = find_and_replace('carried out by her', 'carried out by him')
print(f"  Pronoun her->him: {c} replacements")

# 6. Also check tables
c = 0
for t in doc.tables:
    c += replace_in_table(t, 'Ankita Chetri', 'Abhishek Sharma')
    c += replace_in_table(t, 'UT-231-049-0012', 'UT-231-049-0004')
    c += replace_in_table(t, '23084898', '23084891')
    c += replace_in_table(t, 'Urbimala Hazarika', 'Manas Chakraborty')
print(f"  Table replacements: {c}")

# Save
doc.save(out)
print(f"\nSaved to: {out}")

# Copy to final location
final = r'C:\xampp\htdocs\ICC_Companion\temp\ppt_and_report\my_report\PROJECT_REPORT.docx'
try:
    shutil.copy2(out, final)
    print(f"Copied to: {final}")
    print(f"File size: {os.path.getsize(final) / 1024:.1f} KB")
except PermissionError:
    print(f"Permission denied copying to {final}")
    print(f"File is at: {out}")
