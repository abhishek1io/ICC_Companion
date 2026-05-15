from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

doc = Document()

# Title
title = doc.add_heading('ICC COMPANION', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph('PROJECT REPORT')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].bold = True
subtitle.runs[0].font.size = Pt(18)

doc.add_paragraph('A Comprehensive College Management System')
doc.add_paragraph('')

# Project Overview
doc.add_heading('1. PROJECT OVERVIEW', level=1)
doc.add_paragraph(
    'ICC Companion (Icon Commerce College Campus Portal) is a professional web-based '
    'management system designed to digitize and simplify campus life for both students and staff. '
    'The primary goal is to provide real-time access to academic data, attendance, and campus-wide communications.'
)
doc.add_paragraph('Tagline: "Your all-in-one college management solution"', style='Quote')

# Tech Stack
doc.add_heading('2. TECH STACK', level=1)
tech_table = doc.add_table(rows=5, cols=2)
tech_table.style = 'Table Grid'
tech_data = [
    ('Frontend', 'HTML5, CSS3, Vanilla JavaScript'),
    ('Backend', 'PHP (Server-side logic)'),
    ('Database', 'MySQL (Relational data storage)'),
    ('Server', 'Apache via XAMPP'),
    ('Data Exchange', 'JSON API')
]
for i, (key, val) in enumerate(tech_data):
    tech_table.rows[i].cells[0].text = key
    tech_table.rows[i].cells[1].text = val

# Key Features
doc.add_heading('3. KEY FEATURES', level=1)

doc.add_heading('3.1 Student Dashboard', level=2)
features_student = [
    'Attendance Tracker with visual progress bars',
    'Exam Timetable with countdown timers',
    'Centralized Announcements feed',
    'Lost & Found Portal',
    'Academic Resources access'
]
for f in features_student:
    doc.add_paragraph(f, style='List Bullet')

doc.add_heading('3.2 Admin Management Panel', level=2)
features_admin = [
    'Student Information System (Full CRUD)',
    'Smart Attendance Marker',
    'Exam Controller',
    'Global Communications',
    'Lost & Found Manager',
    'Dynamic File Manager'
]
for f in features_admin:
    doc.add_paragraph(f, style='List Bullet')

# Database Schema
doc.add_heading('4. DATABASE SCHEMA', level=1)
doc.add_paragraph('The system uses MySQL database named "campus_portal" with the following core tables:')

db_table = doc.add_table(rows=10, cols=3)
db_table.style = 'Table Grid'
headers = ['Table Name', 'Description', 'Key Fields']
for i, h in enumerate(headers):
    db_table.rows[0].cells[i].text = h

db_data = [
    ('departments', 'Academic divisions', 'dept_id, dept_code'),
    ('subjects', 'Courses per dept/sem', 'subject_id, dept_code'),
    ('students', 'Student profiles', 'student_id, roll_number'),
    ('admins', 'Staff accounts', 'admin_id, username'),
    ('attendance', 'Daily records', 'attendance_id, roll_number'),
    ('exams', 'Exam schedules', 'exam_id, subject_id'),
    ('lost_found', 'Lost & Found items', 'item_id, item_type'),
    ('announcements', 'System notices', 'announcement_id'),
    ('class_routines', 'Weekly schedules', 'routine_id')
]
for i, row in enumerate(db_data, 1):
    for j, val in enumerate(row):
        db_table.rows[i].cells[j].text = val

# System Architecture
doc.add_heading('5. SYSTEM ARCHITECTURE', level=1)
doc.add_paragraph('The project follows a 3-tier architecture:')
doc.add_paragraph('1. User Interface (HTML/CSS/JS)', style='List Number')
doc.add_paragraph('2. API Layer (PHP Endpoints)', style='List Number')
doc.add_paragraph('3. Data Layer (MySQL Database)', style='List Number')

# Technical Highlights
doc.add_heading('6. TECHNICAL HIGHLIGHTS', level=1)

doc.add_heading('6.1 Database Transactions (ACID)', level=2)
doc.add_paragraph(
    'Used in attendance system for atomic operations. Either all records save or none, '
    'preventing "half-saved" states.'
)

doc.add_heading('6.2 SQL Injection Prevention', level=2)
doc.add_paragraph(
    'All user data handled via Prepared Statements (mysqli_prepare). Industry standard for security.'
)

doc.add_heading('6.3 Asynchronous Data Handling', level=2)
doc.add_paragraph(
    'AJAX via Fetch API and JSON. Updates only necessary data blocks without page refresh.'
)

# Development Roadmap
doc.add_heading('7. DEVELOPMENT ROADMAP', level=1)
roadmap = [
    ('Week 1', 'Foundation - Core architecture, Database setup, Authentication'),
    ('Week 2', 'Student Experience - Attendance, Timetable, Lost & Found'),
    ('Week 3', 'Admin Control - Management tools for students, attendance, exams'),
    ('Week 4', 'Polish & Deployment - Announcements, UI/UX refinements')
]
for week, work in roadmap:
    p = doc.add_paragraph()
    p.add_run(f'{week}: ').bold = True
    p.add_run(work)

# Login Credentials
doc.add_heading('8. LOGIN CREDENTIALS', level=1)
cred_table = doc.add_table(rows=3, cols=3)
cred_table.style = 'Table Grid'
cred_table.rows[0].cells[0].text = 'Portal'
cred_table.rows[0].cells[1].text = 'Username'
cred_table.rows[0].cells[2].text = 'Password'
cred_table.rows[1].cells[0].text = 'Student'
cred_table.rows[1].cells[1].text = 'Roll Number'
cred_table.rows[1].cells[2].text = 'Date of Birth'
cred_table.rows[2].cells[0].text = 'Admin'
cred_table.rows[2].cells[1].text = 'admin / staff'
cred_table.rows[2].cells[2].text = 'admin123 / staff123'

# Footer
doc.add_paragraph('')
doc.add_paragraph('---')
doc.add_paragraph('Report Generated: ICC Companion Project')

output_path = r'C:\xampp\htdocs\ICC_Companion\temp\ppt_and_report\my_report\PROJECT_REPORT.docx'
doc.save(output_path)
print(f'Report saved to: {output_path}')