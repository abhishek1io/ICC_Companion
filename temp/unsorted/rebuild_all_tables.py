from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path
import shutil

DOCPATH = r'C:\xampp\htdocs\ICC_Companion\temp\ppt_and_report\my_report\ICC_Companion_v5.docx'

def make_cell(text, is_header=False, col_width=2500):
    tc = OxmlElement('w:tc')
    tcPr = OxmlElement('w:tcPr')
    tcW = OxmlElement('w:tcW')
    tcW.set(qn('w:w'), str(col_width))
    tcW.set(qn('w:type'), 'dxa')
    tcPr.append(tcW)
    if is_header:
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), 'D9E2F3')
        shading.set(qn('w:val'), 'clear')
        tcPr.append(shading)
    vAlign = OxmlElement('w:vAlign')
    vAlign.set(qn('w:val'), 'center')
    tcPr.append(vAlign)
    tc.append(tcPr)
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    pJc = OxmlElement('w:jc')
    pJc.set(qn('w:val'), 'center')
    pPr.append(pJc)
    p.append(pPr)
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rPr.append(rFonts)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '22')
    rPr.append(sz)
    szCs = OxmlElement('w:szCs')
    szCs.set(qn('w:val'), '22')
    rPr.append(szCs)
    if is_header:
        b = OxmlElement('w:b')
        rPr.append(b)
    r.append(rPr)
    t = OxmlElement('w:t')
    t.text = str(text)
    t.set(qn('xml:space'), 'preserve')
    r.append(t)
    p.append(r)
    tc.append(p)
    return tc

def make_db_table(headers, data_rows):
    num_cols = len(headers)
    tbl = OxmlElement('w:tbl')
    tblPr = OxmlElement('w:tblPr')
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), '9000')
    tblW.set(qn('w:type'), 'dxa')
    tblPr.append(tblW)
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'center')
    tblPr.append(jc)
    tblBorders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{side}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), 'auto')
        tblBorders.append(border)
    tblPr.append(tblBorders)
    tblLayout = OxmlElement('w:tblLayout')
    tblLayout.set(qn('w:type'), 'fixed')
    tblPr.append(tblLayout)
    tbl.append(tblPr)
    tblGrid = OxmlElement('w:tblGrid')
    for ci in range(num_cols):
        gridCol = OxmlElement('w:gridCol')
        cw = [2000, 2000, 2000, 3000][ci] if ci < 4 else 2500
        gridCol.set(qn('w:w'), str(cw))
        tblGrid.append(gridCol)
    tbl.append(tblGrid)
    tr = OxmlElement('w:tr')
    for ci, h in enumerate(headers):
        cw = [2000, 2000, 2000, 3000][ci] if ci < 4 else 2500
        tr.append(make_cell(h, is_header=True, col_width=cw))
    tbl.append(tr)
    for row_data in data_rows:
        tr = OxmlElement('w:tr')
        for ci, val in enumerate(row_data):
            cw = [2000, 2000, 2000, 3000][ci] if ci < 4 else 2500
            tr.append(make_cell(val, col_width=cw))
        tbl.append(tr)
    return tbl

def make_para(text, bold=True, size=12):
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    pJc = OxmlElement('w:jc')
    pJc.set(qn('w:val'), 'center')
    pPr.append(pJc)
    p.append(pPr)
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rPr.append(rFonts)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(size * 2))
    rPr.append(sz)
    if bold:
        b = OxmlElement('w:b')
        rPr.append(b)
    r.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    r.append(t)
    p.append(r)
    return p

def make_spacer():
    sp = OxmlElement('w:p')
    spPr = OxmlElement('w:pPr')
    sp.append(spPr)
    return sp

HEADERS = ["Field Name", "Data Type", "Constraint", "Description"]

ALL_TABLES = [
    ("Table 1: departments (Department Master Data)", [
        ["dept_id", "INT(11)", "PRIMARY KEY, NOT NULL, AUTO_INCREMENT", "Unique identifier for each department"],
        ["dept_code", "VARCHAR(10)", "UNIQUE, NOT NULL", "Short code identifying department (BCA, BBA, BA, BCOM)"],
        ["dept_name", "VARCHAR(100)", "NOT NULL", "Full name of the department"],
        ["max_semesters", "INT(11)", "DEFAULT 6", "Maximum number of semesters in the program"],
    ]),
    ("Table 2: subjects (Subject Catalog)", [
        ["subject_id", "INT(11)", "PRIMARY KEY, NOT NULL, AUTO_INCREMENT", "Unique identifier for each subject"],
        ["subject_code", "VARCHAR(20)", "NOT NULL", "Course code (e.g. BCA605, BBA601)"],
        ["subject_name", "VARCHAR(100)", "NOT NULL", "Full name of the subject"],
        ["dept_code", "VARCHAR(10)", "FOREIGN KEY (dept_code) REFERENCES departments(dept_code)", "Department offering the subject"],
        ["semester", "INT(11)", "NOT NULL", "Semester in which the subject is taught"],
    ]),
    ("Table 3: students (Student Records)", [
        ["student_id", "INT(11)", "PRIMARY KEY, NOT NULL, AUTO_INCREMENT", "Unique identifier for each student"],
        ["roll_number", "VARCHAR(20)", "UNIQUE, NOT NULL", "Institutional roll number assigned to student"],
        ["name", "VARCHAR(100)", "NOT NULL", "Full name of the student"],
        ["dob", "DATE", "NOT NULL", "Date of birth of the student"],
        ["email", "VARCHAR(100)", "NULLABLE", "Email address of the student"],
        ["phone", "VARCHAR(15)", "NULLABLE", "Contact phone number"],
        ["dept_code", "VARCHAR(10)", "FOREIGN KEY (dept_code) REFERENCES departments(dept_code)", "Department the student belongs to"],
        ["semester", "INT(11)", "NOT NULL", "Current semester of the student"],
        ["created_at", "TIMESTAMP", "DEFAULT CURRENT_TIMESTAMP", "Timestamp when the record was created"],
    ]),
    ("Table 4: admins (Administrative System Credentials)", [
        ["admin_id", "INT(11)", "PRIMARY KEY, NOT NULL, AUTO_INCREMENT", "Unique identifier for each admin user"],
        ["username", "VARCHAR(50)", "UNIQUE, NOT NULL", "Login username for the admin panel"],
        ["password", "VARCHAR(255)", "NOT NULL", "Hashed password for authentication"],
        ["name", "VARCHAR(100)", "NOT NULL", "Display name of the admin user"],
        ["role", "ENUM('super-admin','dept-admin','faculty')", "DEFAULT 'faculty'", "Role-based access level"],
        ["assigned_dept", "VARCHAR(10)", "DEFAULT 'all'", "Department scope for dept-admin/faculty roles"],
        ["assigned_semester", "VARCHAR(10)", "DEFAULT 'all'", "Semester scope for faculty role"],
        ["assigned_subject", "VARCHAR(50)", "DEFAULT 'all'", "Subject scope for faculty role"],
    ]),
    ("Table 5: announcements (Institutional Announcements Board)", [
        ["announcement_id", "INT(11)", "PRIMARY KEY, NOT NULL, AUTO_INCREMENT", "Unique identifier for each announcement"],
        ["title", "VARCHAR(200)", "NOT NULL", "Headline title of the announcement"],
        ["description", "TEXT", "NOT NULL", "Full body content of the announcement"],
        ["attachment_url", "VARCHAR(255)", "NULLABLE", "URL to optional file attachment"],
        ["link_url", "VARCHAR(255)", "NULLABLE", "External link URL for reference"],
        ["priority", "ENUM('high','medium','low')", "DEFAULT 'medium'", "Priority level of the announcement"],
        ["target_dept", "VARCHAR(50)", "DEFAULT 'all'", "Target department filter"],
        ["target_semester", "VARCHAR(20)", "DEFAULT 'all'", "Target semester filter"],
        ["posted_by", "VARCHAR(50)", "NULLABLE", "Admin username who posted the announcement"],
        ["created_at", "TIMESTAMP", "DEFAULT CURRENT_TIMESTAMP", "Timestamp when the announcement was posted"],
    ]),
    ("Table 6: attendance (Student Attendance Tracking)", [
        ["attendance_id", "INT(11)", "PRIMARY KEY, NOT NULL, AUTO_INCREMENT", "Unique identifier for each attendance record"],
        ["roll_number", "VARCHAR(20)", "FOREIGN KEY (roll_number) REFERENCES students(roll_number)", "Student roll number being marked"],
        ["subject_id", "INT(11)", "FOREIGN KEY (subject_id) REFERENCES subjects(subject_id)", "Subject for which attendance is marked"],
        ["attendance_date", "DATE", "NOT NULL, UNIQUE KEY (roll_number, subject_id, attendance_date)", "Date of the class"],
        ["status", "ENUM('present','absent')", "DEFAULT 'present'", "Attendance status for the student"],
        ["marked_by", "VARCHAR(50)", "NULLABLE", "Admin/faculty who marked the attendance"],
        ["created_at", "TIMESTAMP", "DEFAULT CURRENT_TIMESTAMP", "Timestamp when record was created"],
    ]),
    ("Table 7: class_routines (Weekly Class Schedule Files)", [
        ["routine_id", "INT(11)", "PRIMARY KEY, NOT NULL, AUTO_INCREMENT", "Unique identifier for each routine entry"],
        ["title", "VARCHAR(200)", "NOT NULL", "Title or description of the routine file"],
        ["file_url", "VARCHAR(255)", "NOT NULL", "URL or path to the uploaded routine PDF"],
        ["dept_code", "VARCHAR(10)", "NOT NULL", "Department code to which the routine applies"],
        ["semester", "VARCHAR(10)", "NOT NULL", "Semester for which the routine applies"],
        ["created_at", "TIMESTAMP", "DEFAULT CURRENT_TIMESTAMP", "Timestamp when routine was uploaded"],
    ]),
    ("Table 8: exams (Exam Scheduling Database)", [
        ["exam_id", "INT(11)", "PRIMARY KEY, NOT NULL, AUTO_INCREMENT", "Unique identifier for each exam entry"],
        ["subject_id", "INT(11)", "FOREIGN KEY (subject_id) REFERENCES subjects(subject_id)", "Subject for which the exam is scheduled"],
        ["exam_date", "DATE", "NOT NULL", "Date of the examination"],
        ["start_time", "TIME", "NOT NULL", "Start time of the examination"],
        ["end_time", "TIME", "NOT NULL", "End time of the examination"],
        ["room", "VARCHAR(50)", "NULLABLE", "Room or hall where the exam is conducted"],
        ["exam_type", "ENUM('sessional','final')", "DEFAULT 'final'", "Type of examination"],
        ["attachment_url", "VARCHAR(255)", "NULLABLE", "URL to optional exam attachment or timetable"],
        ["created_at", "TIMESTAMP", "DEFAULT CURRENT_TIMESTAMP", "Timestamp when exam record was created"],
    ]),
    ("Table 9: exam_schedules (Academic Calendar Scheduling)", [
        ["schedule_id", "INT(11)", "PRIMARY KEY, NOT NULL, AUTO_INCREMENT", "Unique identifier for each schedule file"],
        ["schedule_type", "ENUM('sessional','final')", "NOT NULL", "Type of exam schedule"],
        ["file_url", "VARCHAR(255)", "NOT NULL", "URL or path to the uploaded schedule PDF"],
        ["created_at", "TIMESTAMP", "DEFAULT CURRENT_TIMESTAMP", "Timestamp when schedule was uploaded"],
        ["dept_code", "VARCHAR(10)", "FOREIGN KEY (dept_code) REFERENCES departments(dept_code)", "Department to which the schedule applies"],
        ["semester", "INT(11)", "NULLABLE", "Semester to which the schedule applies"],
    ]),
    ("Table 10: lost_found (Lost and Found Items Registry)", [
        ["item_id", "INT(11)", "PRIMARY KEY, NOT NULL, AUTO_INCREMENT", "Unique identifier for each item entry"],
        ["title", "VARCHAR(100)", "NOT NULL", "Short title describing the item"],
        ["description", "TEXT", "NULLABLE", "Detailed description of the item"],
        ["category", "ENUM('id-card','phone','wallet','books','electronics','other')", "DEFAULT 'other'", "Category classification of the item"],
        ["item_type", "ENUM('lost','found')", "NOT NULL", "Whether the item is lost or found"],
        ["location", "VARCHAR(100)", "NULLABLE", "Location where item was lost or found"],
        ["item_date", "DATE", "NULLABLE", "Date when item was lost or found"],
        ["contact_info", "VARCHAR(100)", "NULLABLE", "Contact information for the item"],
        ["image_url", "VARCHAR(255)", "NULLABLE", "URL to uploaded item image"],
        ["status", "ENUM('active','claimed')", "DEFAULT 'active'", "Current claim status of the item"],
        ["posted_by", "VARCHAR(50)", "NULLABLE", "Admin or student who posted the entry"],
        ["claimed_by", "VARCHAR(20)", "NULLABLE", "Roll number of the claiming student"],
        ["claimed_date", "DATE", "NULLABLE", "Date when the item was claimed"],
        ["created_at", "TIMESTAMP", "DEFAULT CURRENT_TIMESTAMP", "Timestamp when the entry was created"],
    ]),
    ("Table 11: resources (Academic Learning Materials Repository)", [
        ["resource_id", "INT(11)", "PRIMARY KEY, NOT NULL, AUTO_INCREMENT", "Unique identifier for each resource"],
        ["title", "VARCHAR(200)", "NOT NULL", "Title of the learning resource"],
        ["description", "TEXT", "NULLABLE", "Brief description of the resource content"],
        ["resource_type", "ENUM('syllabus','book','material','link','text','others')", "NOT NULL", "Type classification of the resource"],
        ["file_url", "VARCHAR(255)", "NULLABLE", "URL or path to uploaded file resource"],
        ["link_url", "TEXT", "NULLABLE", "External URL for link-type resources"],
        ["content_text", "TEXT", "NULLABLE", "Inline text content for text-type resources"],
        ["dept_code", "VARCHAR(10)", "FOREIGN KEY (dept_code) REFERENCES departments(dept_code)", "Department for which resource is relevant"],
        ["semester", "VARCHAR(10)", "NOT NULL", "Semester for which resource is relevant"],
        ["subject_id", "INT(11)", "FOREIGN KEY (subject_id) REFERENCES subjects(subject_id)", "Subject to which the resource belongs"],
        ["posted_by", "VARCHAR(50)", "NULLABLE", "Admin username who uploaded the resource"],
        ["created_at", "TIMESTAMP", "DEFAULT CURRENT_TIMESTAMP", "Timestamp when resource was uploaded"],
    ]),
    ("Table 12: faculty_subjects (Faculty Subject Assignments)", [
        ["assignment_id", "INT(11)", "PRIMARY KEY, NOT NULL, AUTO_INCREMENT", "Unique identifier for each assignment"],
        ["admin_id", "INT(11)", "FOREIGN KEY (admin_id) REFERENCES admins(admin_id) ON DELETE CASCADE", "Admin or faculty assigned to the subject"],
        ["subject_id", "INT(11)", "FOREIGN KEY (subject_id) REFERENCES subjects(subject_id) ON DELETE CASCADE", "Subject assigned to the faculty member"],
    ]),
]

print("Loading document...")
doc = Document(DOCPATH)
body = doc.element.body

# STEP 1: Find intro paragraph and CHAPTER 8
intro_elem = None
ch8_elem = None
first_heading = None

for elem in body:
    tag = elem.tag.split('}')[-1]
    if tag == 'p':
        texts = [t.text for t in elem.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t') if t.text]
        txt = ''.join(texts).strip()
        if "The application database, named 'campus_portal'" in txt:
            intro_elem = elem
        if txt == 'CHAPTER 8':
            ch8_elem = elem

if not intro_elem or not ch8_elem:
    print("ERROR: Could not find intro paragraph or CHAPTER 8")
    exit(1)

print(f"Found intro paragraph and CHAPTER 8")

# STEP 2: Remove all elements between intro and CHAPTER 8
to_remove = []
found_intro = False
for elem in body:
    tag = elem.tag.split('}')[-1]
    if elem is intro_elem:
        found_intro = True
        continue
    if elem is ch8_elem:
        break
    if found_intro:
        to_remove.append(elem)

print(f"Removing {len(to_remove)} old elements...")
for elem in reversed(to_remove):
    body.remove(elem)

# STEP 3: Insert all 12 table blocks after intro
# addnext inserts right after ref_elem.
# To achieve: ref → T1 → sp → T1_tbl → sp → T2 → sp → T2_tbl → ... → ch8
# We must insert in REVERSE order because each addnext pushes subsequent items farther away.
# So faculty_subjects (last table) is added first, then resources (second-to-last), etc.
# Within each block: heading, sp1, table, sp2
# For a single block to appear as heading→sp1→tbl→sp2 after ref:
# addnext(sp2) first, then addnext(tbl), then addnext(sp1), then addnext(heading)

all_block_elems = []
for heading_text, rows_data in ALL_TABLES:
    heading = make_para(heading_text, bold=True, size=12)
    sp1 = make_spacer()
    tbl = make_db_table(HEADERS, rows_data)
    sp2 = make_spacer()
    all_block_elems.append([heading, sp1, tbl, sp2])

# Insert in reverse order (faculty_subjects first, farthest from intro)
# Within each block, insert in reverse order (sp2 first, heading last, closest to intro)
for block in reversed(all_block_elems):
    for elem in reversed(block):
        intro_elem.addnext(elem)

# STEP 4: Fix text mentions
print("Fixing text mentions...")
for p in doc.paragraphs:
    for r in p.runs:
        if 'eleven well-structured' in r.text:
            r.text = r.text.replace('eleven well-structured', '12 well-structured')

# Add VIEW mention
for p in doc.paragraphs:
    for r in p.runs:
        if '12 normalized MySQL tables' in r.text and 'student_attendance_summary' not in r.text:
            r.text = r.text.replace('12 normalized MySQL tables', '12 normalized MySQL tables and a student_attendance_summary view')
        if '12 well-structured tables' in r.text and 'student_attendance_summary' not in r.text:
            pass  # already has full text

# Save to temp then copy
tmp_path = Path(DOCPATH).parent / "ICC_Companion_v5_tmp.docx"
doc.save(str(tmp_path))
import time
time.sleep(0.5)
shutil.copy2(str(tmp_path), DOCPATH)
tmp_path.unlink()
print(f"Saved to {DOCPATH}")

# VERIFY
print("\n=== VERIFICATION ===")
doc2 = Document(DOCPATH)
print(f"Total python-docx tables: {len(doc2.tables)}")
ft_count = 0
for t in doc2.tables:
    if t.rows and t.rows[0].cells[0].text.strip() == 'Field Name':
        ft_count += 1
print(f"DB schema tables: {ft_count}")

print("\nDocument order:")
for elem in doc2.element.body:
    tag = elem.tag.split('}')[-1]
    if tag == 'p':
        texts = [t.text for t in elem.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t') if t.text]
        txt = ''.join(texts).strip()
        if txt.startswith('Table ') and ':' in txt and len(txt) < 150:
            print(f"  {txt}")
    elif tag == 'tbl':
        rows = elem.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr')
        if rows:
            first = rows[0].find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')
            if first is not None and first.text and first.text.strip() == 'Field Name' and len(rows) > 1:
                second = rows[1].find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')
                name = second.text if second is not None else '?'
                print(f"  -> [{name}] ({len(rows)} fields)")

print("\nDONE!")
