from docx import Document
doc = Document(r'C:\xampp\htdocs\ICC_Companion\temp\ppt_and_report\my_report\ICC_Companion_v5.docx')

for p in doc.paragraphs:
    if 'twelve well-structured tables' in p.text:
        for run in p.runs:
            if 'twelve well-structured tables' in run.text:
                run.text = run.text.replace(
                    'twelve well-structured tables',
                    '12 tables and 1 database view (student_attendance_summary)'
                )
                print('Fixed: added view mention')
                break
        break

doc.save(r'C:\xampp\htdocs\ICC_Companion\temp\ppt_and_report\my_report\ICC_Companion_v5.docx')
print('Saved')

# verify
doc2 = Document(r'C:\xampp\htdocs\ICC_Companion\temp\ppt_and_report\my_report\ICC_Companion_v5.docx')
for p in doc2.paragraphs:
    if '12 tables' in p.text:
        print(f'Verified: "{p.text[:100]}..."')
