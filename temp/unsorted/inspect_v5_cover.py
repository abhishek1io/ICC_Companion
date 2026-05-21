from docx import Document
doc = Document(r'C:\xampp\htdocs\ICC_Companion\temp\ppt_and_report\my_report\ICC_Companion_v5.docx')
# Look at cover page text
for i in range(20):
    p = doc.paragraphs[i]
    txt = p.text.strip()
    if txt:
        print(f'P[{i}]: "{txt[:150]}"')
print('\n--- Cover table info ---')
t0 = doc.tables[0]
for ri, row in enumerate(t0.rows):
    for ci, cell in enumerate(row.cells):
        txt = cell.text.strip()[:80]
        if txt:
            print(f'  Table 0, R{ri}C{ci}: "{txt}"')
