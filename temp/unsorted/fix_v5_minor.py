from docx import Document
doc = Document(r'C:\xampp\htdocs\ICC_Companion\temp\ppt_and_report\my_report\ICC_Companion_v5.docx')

# Fix "eleven" -> "twelve"
for p in doc.paragraphs:
    if 'eleven' in p.text:
        for run in p.runs:
            if 'eleven' in run.text:
                run.text = run.text.replace('eleven', 'twelve')
                print(f'Fixed: "{run.text[:60]}..."')
                break

# Fix Table 7 <-> Table 8 swap
# Table 7 should be exam_schedules, Table 8 should be faculty_subjects
for p in doc.paragraphs:
    txt = p.text.strip()
    if txt.startswith('Table 7:') and 'exam_schedules' not in txt:
        # This is actually Table 8 content with wrong number
        for run in p.runs:
            if 'Table 7:' in run.text:
                run.text = run.text.replace('Table 7:', 'Table 8:')
                print(f'Fixed heading: {run.text}')
    elif txt.startswith('Table 8:') and 'faculty_subjects' not in txt:
        for run in p.runs:
            if 'Table 8:' in run.text:
                run.text = run.text.replace('Table 8:', 'Table 7:')
                print(f'Fixed heading: {run.text}')

doc.save(r'C:\xampp\htdocs\ICC_Companion\temp\ppt_and_report\my_report\ICC_Companion_v5.docx')
print('Done')

# Verify
doc2 = Document(r'C:\xampp\htdocs\ICC_Companion\temp\ppt_and_report\my_report\ICC_Companion_v5.docx')
for p in doc2.paragraphs:
    txt = p.text.strip()
    if 'eleven' in txt:
        print(f'STILL ELEVEN: {txt}')
    if txt.startswith('Table 7:') or txt.startswith('Table 8:'):
        print(f'Heading: {txt}')
