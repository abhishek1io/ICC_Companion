from docx import Document
doc = Document(r'C:\xampp\htdocs\ICC_Companion\temp\ppt_and_report\my_report\ICC_Companion_v5.docx')
print(f'Total paragraphs: {len(doc.paragraphs)}')
print(f'Total tables: {len(doc.tables)}')
print()
for i, t in enumerate(doc.tables):
    rows = len(t.rows)
    cols = len(t.columns)
    first = t.rows[0].cells[0].text[:80] if rows > 0 else 'empty'
    print(f'Table {i}: {rows}r x {cols}c | "{first}"')
print()
print('Paragraphs containing table references:')
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith('Table ') and len(txt) < 60:
        print(f'  P[{i}]: "{txt}"')
