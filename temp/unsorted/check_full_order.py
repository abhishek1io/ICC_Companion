from docx import Document
doc = Document(r'C:\xampp\htdocs\ICC_Companion\temp\ppt_and_report\my_report\ICC_Companion_v5.docx')

# Print ALL elements between "Table 6: lost_found" and "CHAPTER 8"
print('Elements from lost_found heading to CHAPTER 8:')
in_section = False
seen_8 = False
for elem in doc.element.body:
    tag = elem.tag.split('}')[-1]
    if tag == 'p':
        texts = [t.text for t in elem.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t') if t.text]
        txt = ''.join(texts).strip()
        if 'CHAPTER 8' in txt:
            print(f'  >>> CHAPTER 8 FOUND <<<')
            break
        if 'Table 6:' in txt:
            in_section = True
        if in_section:
            if txt:
                print(f'  p: "{txt[:80]}"')
            else:
                print(f'  p: (empty)')
    elif tag == 'tbl' and in_section:
        try:
            t_elem = elem.find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')
            first_text = t_elem.text if t_elem is not None else 'unknown'
            # get second row to identify the table
            rows = elem.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr')
            second_text = ''
            if len(rows) > 1:
                tc2 = rows[1].find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')
                if tc2 is not None and tc2.text:
                    second_text = tc2.text[:25]
            print(f'  tbl: "{first_text[:40]}" -> "{second_text}"')
        except:
            print(f'  tbl: (error)')

print(f'\nTotal tables: {len(doc.tables)}')

# Show all db tables with their names
print('\nAll DB tables in order:')
db_i = 0
for i, t in enumerate(doc.tables):
    if t.rows and t.rows[0].cells[0].text.strip() == 'Field Name':
        db_i += 1
        name = t.rows[1].cells[0].text.strip()[:30] if len(t.rows) > 1 else '?'
        print(f'  DB#{db_i} = Table[{i}]: "{name}" ({len(t.rows)} rows)')
