from docx import Document
doc = Document(r'C:\xampp\htdocs\ICC_Companion\temp\ppt_and_report\my_report\ICC_Companion_v5.docx')
print('Cover page text:')
for i in range(15):
    txt = doc.paragraphs[i].text.strip()
    if txt:
        print(f'  P[{i}]: "{txt[:120]}"')
print()
print('Tables:')
for i, t in enumerate(doc.tables):
    rows = len(t.rows)
    first = t.rows[0].cells[0].text.strip()[:60] if rows > 0 else 'empty'
    print(f'  Table {i} ({rows}r): "{first}"')
print()
print('Database section:')
for i in range(295, 330):
    p = doc.paragraphs[i]
    if p.text.strip():
        print(f'  P[{i}]: "{p.text.strip()[:120]}"')
