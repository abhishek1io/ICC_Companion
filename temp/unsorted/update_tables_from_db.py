"""
Update v5 docx tables to match actual DB schema from campus_portal.sql export.
Rebuilds the 6 new tables with correct fields, constraints, descriptions.
"""
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

V5_PATH = r'C:\xampp\htdocs\ICC_Companion\temp\ppt_and_report\my_report\ICC_Companion_v5.docx'
doc = Document(V5_PATH)

def make_table_element(headers, data):
    cols = len(headers)
    tbl = OxmlElement('w:tbl')
    tblPr = OxmlElement('w:tblPr')
    tblStyle = OxmlElement('w:tblStyle')
    tblStyle.set(qn('w:val'), 'Table Grid')
    tblPr.append(tblStyle)
    tblWidth = OxmlElement('w:tblW')
    tblWidth.set(qn('w:w'), '5000')
    tblWidth.set(qn('w:type'), 'pct')
    tblPr.append(tblWidth)
    tbl.append(tblPr)
    tblGrid = OxmlElement('w:tblGrid')
    for _ in range(cols):
        gridCol = OxmlElement('w:gridCol')
        gridCol.set(qn('w:w'), '2500')
        tblGrid.append(gridCol)
    tbl.append(tblGrid)

    def make_cell(text, is_header=False):
        tc = OxmlElement('w:tc')
        tcPr = OxmlElement('w:tcPr')
        tcW = OxmlElement('w:tcW')
        tcW.set(qn('w:w'), '2500')
        tcW.set(qn('w:type'), 'dxa')
        tcPr.append(tcW)
        if is_header:
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), 'D9E2F3')
            shading.set(qn('w:val'), 'clear')
            tcPr.append(shading)
        tc.append(tcPr)
        p = OxmlElement('w:p')
        pPr = OxmlElement('w:pPr')
        pJc = OxmlElement('w:jc')
        pJc.set(qn('w:val'), 'center')
        pPr.append(pJc)
        p.append(pPr)
        r = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), 'Times New Roman')
        rFonts.set(qn('w:hAnsi'), 'Times New Roman')
        rPr.append(rFonts)
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), '22')
        rPr.append(sz)
        if is_header:
            b = OxmlElement('w:b')
            rPr.append(b)
        r.append(rPr)
        t = OxmlElement('w:t')
        t.text = str(text)
        t.set(qn('xml:space'), 'preserve')
        r.append(t)
        p.append(r)
        tc.append(p)
        return tc

    tr = OxmlElement('w:tr')
    for h in headers:
        tr.append(make_cell(h, is_header=True))
    tbl.append(tr)
    for row_data in data:
        tr = OxmlElement('w:tr')
        for val in row_data:
            tr.append(make_cell(val))
        tbl.append(tr)
    return tbl

def rebuild_table(start_marker, headers, data):
    """Find a table by the first data cell marker and replace it."""
    for elem in doc.element.body:
        tag = elem.tag.split('}')[-1]
        if tag == 'tbl':
            try:
                rows = elem.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr')
                if len(rows) > 1:
                    tc = rows[1].find('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')
                    if tc is not None and tc.text and tc.text.strip() == start_marker:
                        # Found it - replace with new table
                        new_tbl = make_table_element(headers, data)
                        elem.getparent().replace(elem, new_tbl)
                        print(f'  Rebuilt table starting with "{start_marker}"')
                        return True
            except:
                pass
    print(f'  WARNING: Table starting with "{start_marker}" not found!')
    return False

# ===== REBUILD announcements table =====
rebuild_table('announcement_id',
    ['Field Name', 'Data Type', 'Constraint', 'Description'],
    [
        ('announcement_id', 'INT(11)', 'PK, AUTO_INCREMENT', 'Unique announcement ID'),
        ('title', 'VARCHAR(200)', 'NOT NULL', 'Announcement title'),
        ('description', 'TEXT', 'NOT NULL', 'Announcement content/body'),
        ('attachment_url', 'VARCHAR(255)', 'NULL', 'Optional file attachment path'),
        ('link_url', 'VARCHAR(255)', 'NULL', 'External link URL'),
        ('priority', "ENUM('high','medium','low')", "DEFAULT 'medium'", 'Priority level'),
        ('target_dept', "VARCHAR(50)", "DEFAULT 'all'", 'Target department or "all"'),
        ('target_semester', "VARCHAR(20)", "DEFAULT 'all'", 'Target semester or "all"'),
        ('posted_by', 'VARCHAR(50)', 'NULL', 'Admin username who posted'),
        ('created_at', 'TIMESTAMP', 'DEFAULT CURRENT_TIMESTAMP', 'Posting timestamp'),
    ]
)

# ===== REBUILD class_routines table =====
rebuild_table('routine_id',
    ['Field Name', 'Data Type', 'Constraint', 'Description'],
    [
        ('routine_id', 'INT(11)', 'PK, AUTO_INCREMENT', 'Unique routine ID'),
        ('title', 'VARCHAR(200)', 'NOT NULL', 'Routine display title'),
        ('file_url', 'VARCHAR(255)', 'NOT NULL', 'PDF/image file path'),
        ('dept_code', 'VARCHAR(10)', 'NOT NULL', 'Department code'),
        ('semester', 'VARCHAR(10)', 'NOT NULL', 'Semester number'),
        ('created_at', 'TIMESTAMP', 'DEFAULT CURRENT_TIMESTAMP', 'Upload timestamp'),
    ]
)

# ===== REBUILD resources table =====
rebuild_table('resource_id',
    ['Field Name', 'Data Type', 'Constraint', 'Description'],
    [
        ('resource_id', 'INT(11)', 'PK, AUTO_INCREMENT', 'Unique resource ID'),
        ('title', 'VARCHAR(200)', 'NOT NULL', 'Resource title'),
        ('description', 'TEXT', 'NULL', 'Resource description'),
        ('resource_type', "ENUM('syllabus','book','material','link','text','others')", 'NOT NULL', 'Type of resource'),
        ('file_url', 'VARCHAR(255)', 'NULL', 'Uploaded file path'),
        ('link_url', 'TEXT', 'NULL', 'External resource URL'),
        ('content_text', 'TEXT', 'NULL', 'Inline text content'),
        ('dept_code', 'VARCHAR(10)', 'NOT NULL', 'Department code'),
        ('semester', 'VARCHAR(10)', 'NOT NULL', 'Semester number'),
        ('subject_id', 'INT(11)', 'FK -> subjects(subject_id)', 'Linked subject (optional)'),
        ('posted_by', 'VARCHAR(50)', 'NULL', 'Admin username who posted'),
        ('created_at', 'TIMESTAMP', 'DEFAULT CURRENT_TIMESTAMP', 'Upload timestamp'),
    ]
)

# ===== REBUILD lost_found table =====
rebuild_table('item_id',
    ['Field Name', 'Data Type', 'Constraint', 'Description'],
    [
        ('item_id', 'INT(11)', 'PK, AUTO_INCREMENT', 'Unique item ID'),
        ('title', 'VARCHAR(100)', 'NOT NULL', 'Item title'),
        ('description', 'TEXT', 'NULL', 'Item description'),
        ('category', "ENUM('id-card','phone','wallet','books','electronics','other')", "DEFAULT 'other'", 'Item category'),
        ('item_type', "ENUM('lost','found')", 'NOT NULL', 'Lost or Found'),
        ('location', 'VARCHAR(100)', 'NULL', 'Where item was lost/found'),
        ('item_date', 'DATE', 'NULL', 'Date item was lost/found'),
        ('contact_info', 'VARCHAR(100)', 'NULL', 'Contact info of reporter'),
        ('image_url', 'VARCHAR(255)', 'NULL', 'Item image path'),
        ('status', "ENUM('active','claimed')", "DEFAULT 'active'", 'Current item status'),
        ('posted_by', 'VARCHAR(50)', 'NULL', 'Username of person who posted'),
        ('claimed_by', 'VARCHAR(20)', 'NULL', 'Roll number of claimant'),
        ('claimed_date', 'DATE', 'NULL', 'Date item was claimed'),
        ('created_at', 'TIMESTAMP', 'DEFAULT CURRENT_TIMESTAMP', 'Submission timestamp'),
    ]
)

# ===== REBUILD exam_schedules table =====
rebuild_table('schedule_id',
    ['Field Name', 'Data Type', 'Constraint', 'Description'],
    [
        ('schedule_id', 'INT(11)', 'PK, AUTO_INCREMENT', 'Unique schedule ID'),
        ('schedule_type', "ENUM('sessional','final')", 'NOT NULL', 'Type of exam schedule'),
        ('file_url', 'VARCHAR(255)', 'NOT NULL', 'Schedule file path (PDF)'),
        ('created_at', 'TIMESTAMP', 'DEFAULT CURRENT_TIMESTAMP', 'Upload timestamp'),
        ('dept_code', 'VARCHAR(10)', 'FK -> departments(dept_code)', 'Department code'),
        ('semester', 'INT(11)', 'NULL', 'Semester number'),
    ]
)

# ===== REBUILD faculty_subjects table =====
rebuild_table('assignment_id',
    ['Field Name', 'Data Type', 'Constraint', 'Description'],
    [
        ('assignment_id', 'INT(11)', 'PK, AUTO_INCREMENT', 'Unique assignment ID'),
        ('admin_id', 'INT(11)', 'FK -> admins(admin_id) ON DELETE CASCADE', 'Faculty/admin assigned'),
        ('subject_id', 'INT(11)', 'FK -> subjects(subject_id) ON DELETE CASCADE', 'Subject assigned'),
    ]
)

# ===== SAVE =====
doc.save(V5_PATH)
print(f'\nSaved: {V5_PATH}')

# Verify
doc2 = Document(V5_PATH)
print(f'Tables in doc: {len(doc2.tables)}')
db_count = 0
for i, t in enumerate(doc2.tables):
    if t.rows and t.rows[0].cells[0].text.strip() == 'Field Name':
        db_count += 1
        name = t.rows[1].cells[0].text.strip()[:30] if len(t.rows) > 1 else '?'
        rows = len(t.rows)
        print(f'  DB Table: {name:25s} ({rows} rows)')
print(f'Database tables: {db_count}')
